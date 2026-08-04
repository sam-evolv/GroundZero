from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/samdonworth/GroundZero/vault/briefs/2026-08-04-renewable-operations-business-proposition.docx"
NAVY = "142C3E"
GOLD = "B8892D"
INK = "24333E"
MUTED = "62717C"
PALE = "EEF3F5"
PALE_GOLD = "F7F1E5"
WHITE = "FFFFFF"

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.52)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(9.1)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.03
for name, size, colour in [("Title", 25, NAVY), ("Heading 1", 15.5, NAVY), ("Heading 2", 11, GOLD)]:
    s = styles[name]
    s.font.name = "Aptos Display"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(colour)
    s.paragraph_format.space_before = Pt(5)
    s.paragraph_format.space_after = Pt(3)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("RENEWABLE OPERATIONS INTELLIGENCE  |  BUSINESS PROPOSITION")
r.font.name = "Aptos"
r.font.size = Pt(7.2)
r.font.color.rgb = RGBColor.from_string(MUTED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Confidential founder strategy  •  4 August 2026  •  ")
r.font.size = Pt(7.2)
r.font.color.rgb = RGBColor.from_string(MUTED)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=90, start=150, bottom=90, end=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def title(text, subtitle):
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(1)
    p.add_run(text)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(subtitle)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(GOLD)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:color"), GOLD)
    pbdr.append(bottom)
    pPr.append(pbdr)


def h1(text): doc.add_heading(text, level=1)
def h2(text): doc.add_heading(text, level=2)


def body(text, bold_prefix=None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def bullets(items, compact=True):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.19)
        p.paragraph_format.first_line_indent = Inches(-0.13)
        p.paragraph_format.space_after = Pt(1.3 if compact else 2.5)
        p.add_run(item)


def callout(kicker, text, fill=PALE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(7.05)
    cell = table.cell(0, 0)
    cell.width = Inches(7.05)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade_cell(cell, fill)
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(kicker.upper())
    r.bold = True
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(GOLD)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11.2)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def page():
    doc.add_page_break()

# PAGE 1

title("Renewable Operations Intelligence", "An Ireland-first managed operations company for renewable asset owners")
callout("Core proposition", "Fewer unexplained losses, better contractor control and a verified operating record for every asset, fault, intervention and recovery.")
h1("The business")
body("The company sits above existing O&M contractors and technology vendors. It connects plant data, field evidence and maintenance activity so owners can identify losses, direct action and prove that problems were solved.")
body("The first customer is an independent owner, developer retaining assets or investment vehicle with roughly 20–250 MW of operational solar. These portfolios are large enough to lose meaningful money through weak operations, but often too small to justify a complete internal asset-management team or enterprise software deployment.")
h2("Why now")
bullets([
    "Ireland had approximately 1.59 GW of utility-scale solar connected by May 2026, within 2.7 GW of total connected solar, and targets 8 GW by 2030.",
    "The adjacent UK market had 22.8 GW of solar by June 2026, with roughly 59% estimated to be ground-mounted.",
    "Solar assets operate for 30 years or more, creating a durable operations and maintenance base.",
    "Thermal drones, satellite imagery, inexpensive sensors, robotic equipment and capable AI now make an integrated managed service economically possible.",
])
h2("The problem")
body("Owners can have SCADA, inverter portals, meters, weather data, contractor tickets, warranty documents, inspection PDFs and spreadsheets without one trusted operating picture.")
bullets([
    "Faults and underperformance are found late or never linked to permanent resolution.",
    "Responsibility is unclear across the owner, EPC, OEM, O&M provider and grid.",
    "Warranty claims stall because serial numbers, readings and photographs are missing.",
    "Cleaning, vegetation and inspections are commissioned without a measured economic case.",
    "Monthly investor and lender reporting remains manual, and completed work is not tied to verified output recovery.",
])
callout("Positioning", "Not a solar cleaner. Not another dashboard. The independent operating intelligence layer between plant data, contractors and performance evidence.", PALE_GOLD)

# PAGE 2
page()
title("The proposition", "A productised operations desk with an integrated field-assurance network")
h1("Three layers")
h2("1. Portfolio operations desk")
bullets([
    "Read-only ingestion from SCADA, inverter portals, meters, weather and market sources.",
    "Verified asset hierarchy, generation/export reconciliation and performance baselines.",
    "Business-hours anomaly review ranked by safety, affected capacity and estimated lost MWh.",
    "Owner-visible work orders, contractor KPIs and source-linked monthly reporting.",
])
h2("2. Field assurance network")
body("Qualified partners perform thermal/RGB inspections, electrical work, cleaning, vegetation, drainage, security and environmental checks. The company controls the operating record, coordinates action and verifies the result. This keeps launch capital and certification requirements low while giving the owner one accountable interface.")
h2("3. Evidence and intelligence layer")
bullets([
    "Signal: monitoring, weather, meter or imagery data identifies a problem.",
    "Diagnosis: estimate urgency, affected capacity, likely cause and value at risk.",
    "Action: route work to the responsible contractor or OEM.",
    "Evidence: capture readings, parts, serial numbers, photographs, cost and completion proof.",
    "Verification: use plant data to confirm recovery and close the issue.",
])
h1("The technology-led USP")
h2("Immediate edge")
bullets([
    "Vendor-neutral owner assurance, rather than software tied to one OEM or contractor.",
    "Measures action and recovery, not just alarms.",
    "Combines operational data, contractor work and field evidence in one workflow.",
    "Recommends cleaning, mowing or inspection only when the expected value supports the cost.",
])
h2("New developments used intelligently")
bullets([
    "AI operations copilot: normalise alarms, rank loss exposure, identify repeat faults, extract evidence and draft source-linked reports.",
    "Multimodal evidence graph: connect SCADA, component records, work orders, warranties, drone imagery and post-repair performance.",
    "Drone-to-work-order automation: convert thermal findings into prioritised defects and verify repairs against production data.",
    "Evidence-triggered robotics: deploy cleaning and mowing robots only on economically and technically suitable sites.",
    "Satellite intelligence: use Planet imagery for broad vegetation, flooding, access and construction screening, not module diagnosis.",
    "Cross-fleet learning: use clean fault and serial-number history to identify component cohorts and future maintenance demand.",
])
callout("Long-term moat", "A compounding asset-to-fault-to-action-to-recovery dataset across multiple renewable portfolios.")

# PAGE 3
page()
title("Commercial model and route to scale", "Managed service first, software and robotics earned through paid operations")
h1("The entry offer")
body("Sell a 60–90 day operational baseline and leakage audit before building software or buying equipment.")
bullets([
    "Connect historical and read-only data.",
    "Establish asset and performance baselines.",
    "Identify unresolved losses, contractor delays and evidence gaps.",
    "Produce an owner-ready operating pack.",
    "Close at least one complete signal-to-action-to-verification loop.",
])
callout("Indicative pilot", "€4,000–€8,000 per site. Recurring test price: €750 monthly portfolio base plus approximately €20 per MW per month.", PALE_GOLD)
h2("Economics")
bullets([
    "Five 30 MW sites at the test formula produce approximately €81,000 ARR.",
    "That is strong validation income, but not yet enough for two full-time founders.",
    "Scale requires larger portfolios, higher-value assurance modules and UK expansion.",
    "Field work is passed through transparently with a coordination and assurance fee.",
])
h2("First 90 days")
bullets([
    "Clear employment, confidentiality, IP and conflict boundaries.",
    "Use three to five warm renewable relationships for confidential problem interviews.",
    "Inspect redacted reporting, alarms, contractor tickets and unresolved discrepancies.",
    "Sell one paid baseline, run it manually with cleanly owned Evolv capabilities, and measure value created.",
    "Convert the pilot to a 12-month service. Build reusable software only after the same requirement appears across three customers.",
])
h2("Expansion path")
bullets([
    "Phase 1: solar monitoring, reconciliation, contractor evidence and reporting.",
    "Phase 2: integrated drone, cleaning, robotic mowing, drainage, security and environmental services.",
    "Phase 3: cross-fleet benchmarking, predictive maintenance and lender/insurer reporting.",
    "Phase 4: wind and storage reporting for existing multi-technology customers, with specialist partners.",
    "Phase 5: Irish deployment or service partnerships for proven construction and maintenance robotics.",
])
h2("Deliberate boundaries")
bullets([
    "No 24/7 control room, HV switching, plant-control responsibility or availability guarantees at launch.",
    "No proprietary robot manufacturing or speculative rooftop drone installation.",
    "No major platform build until paid demand and reusable workflows are proven.",
])
callout("Decision gate", "Proceed only when one portfolio owner pays for the operational baseline and grants read-only data access.")
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(2)
p.paragraph_format.space_after = Pt(0)
r = p.add_run("Evidence base: Solar Ireland Scale of Solar 2026; ESB Networks Q1 2026; SolarPower Europe O&M Best Practice Guidelines v6.0; UK DESNZ solar deployment data; named vendor materials.")
r.italic = True
r.font.size = Pt(7.2)
r.font.color.rgb = RGBColor.from_string(MUTED)

props = doc.core_properties
props.title = "Renewable Operations Intelligence Business Proposition"
props.subject = "Ireland-first renewable asset operations, evidence and emerging technology"
props.author = "Hermes Agent for Sam Donworth"
doc.save(OUT)
print(OUT)
