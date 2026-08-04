from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('/Users/samdonworth/GroundZero/vault/briefs/2026-08-04-best-robotics-opportunities-two-page.docx')

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.02
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_opportunity(doc, number, title, verdict, why, risk):
    p = doc.add_paragraph()
    p.style = doc.styles['Heading 2']
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f'{number}. {title}')
    r.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.02
    p.add_run('Offer: ').bold = True
    p.add_run(verdict)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.02
    p.add_run('Why it fits: ').bold = True
    p.add_run(why)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.02
    p.add_run('Main risk: ').bold = True
    p.add_run(risk)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.58)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.68)
section.right_margin = Inches(0.68)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(9.4)
styles['Normal'].paragraph_format.space_after = Pt(3)
styles['Normal'].paragraph_format.line_spacing = 1.03

styles['Title'].font.name = 'Aptos Display'
styles['Title'].font.size = Pt(25)
styles['Title'].font.bold = True
styles['Title'].font.color.rgb = RGBColor(24, 31, 41)

styles['Heading 1'].font.name = 'Aptos Display'
styles['Heading 1'].font.size = Pt(16)
styles['Heading 1'].font.bold = True
styles['Heading 1'].font.color.rgb = RGBColor(24, 31, 41)
styles['Heading 1'].paragraph_format.space_before = Pt(4)
styles['Heading 1'].paragraph_format.space_after = Pt(5)

styles['Heading 2'].font.name = 'Aptos'
styles['Heading 2'].font.size = Pt(11.5)
styles['Heading 2'].font.bold = True
styles['Heading 2'].font.color.rgb = RGBColor(31, 77, 70)

# Page 1
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.paragraph_format.space_after = Pt(2)
p.add_run('The best robotics opportunities for us')

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(9)
r = p.add_run('Ireland first  |  Part-time validation  |  UK and EU scale')
r.bold = True
r.font.color.rgb = RGBColor(90, 98, 108)

callout = doc.add_table(rows=1, cols=1)
callout.autofit = True
cell = callout.cell(0, 0)
set_cell_shading(cell, 'EAF2EF')
set_cell_margins(cell, top=130, start=170, bottom=130, end=170)
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(0)
p.add_run('Best first move: ').bold = True
p.add_run('sell a digital plant-rounds assessment before buying a robot. Prove that customers value the route evidence, then automate the capture.')

doc.add_paragraph('')

add_opportunity(
    doc, '1', 'Digital plant and utility rounds',
    'Map one existing non-GMP route and capture structured RGB, thermal and acoustic evidence. Record gauges, indicators and exceptions, with every decision reviewed by a human.',
    'It combines plant-engineering access with software and product skills. It can start without hardware, creates recurring asset data and can later use wheeled robots, quadrupeds or fixed sensors.',
    'Korial, ANYbotics and other global vendors already serve industrial inspection. We only win by owning a very specific Irish route, evidence requirement and deployment service.'
)

add_opportunity(
    doc, '2', 'No-entry water-asset cleaning',
    'Use established remote equipment to clean sediment from compatible non-product-contact tanks, cooling-tower basins or reservoirs while producing before-and-after evidence.',
    'The value is concrete: less confined-space entry, shutdown, drain/refill work and water loss. Customers can buy it as a scheduled service.',
    'It is operationally heavy. Transport, pumps, hoses, disposal, insurance, recovery planning and backup equipment make it harder for two part-time founders.'
)

add_opportunity(
    doc, '3', 'Controlled-area UV-C decontamination',
    'Qualify and operate an established autonomous UV-C platform in Grade C/D or controlled non-classified pharma and medtech areas.',
    'Ireland has concentrated pharma demand, regulatory pressure for contamination control and commercial robots with digital cycle records.',
    'Quality approval, validation and procurement are demanding. Existing suppliers already operate in this category, so a real service gap must be proven.'
)

add_opportunity(
    doc, '4', 'Private drainage and stormwater records',
    'Provide structured crawler or push-camera surveys for developments, industrial estates and private infrastructure, replacing unindexed video with usable defect and handover records.',
    'This is the fastest route through the property network, requires cheaper equipment and can be scheduled project by project.',
    'Inspection is competitive and can remain a local service. The defensibility must come from longitudinal asset records and workflow integration.'
)

# Page 2
doc.add_page_break()

doc.add_heading('What we should do now', level=1)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(5)
p.add_run('Focus on opportunity 1 for 30 days. ').bold = True
p.add_run('Keep the other three as named pivots. Do not research, build or buy for all four at once.')

steps = [
    ('1. Interview 15 buyers and users.', 'Ask reliability, maintenance, utilities and FM people about the last completed route, not whether they like robots.'),
    ('2. Find one repeated route.', 'Look for the same observable checks at three sites, with a named budget owner and measurable annual burden.'),
    ('3. Complete three to five approved site walks.', 'Choose a simple non-GMP route with 10 to 20 checkpoints, no ATEX classification and no process-control interaction.'),
    ('4. Produce one manual evidence pack.', 'Use customer-owned or borrowed RGB, thermal and acoustic tools. Retain raw evidence and require human review.'),
    ('5. Ask for payment.', 'Sell a bounded assessment or concierge pilot before building bespoke software or integrating mobility.'),
]
for lead, rest in steps:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.02
    p.add_run(lead + ' ').bold = True
    p.add_run(rest)

callout2 = doc.add_table(rows=1, cols=1)
cell = callout2.cell(0, 0)
set_cell_shading(cell, 'FFF3D8')
set_cell_margins(cell, top=125, start=170, bottom=125, end=170)
p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(0)
p.add_run('Hard commercial gate: ').bold = True
p.add_run('continue only if three sites share the route and at least one buyer funds an assessment or pilot. Interest without payment is not validation.')

doc.add_heading('The first pilot should be deliberately boring', level=1)
for text in [
    'One site and one existing route.',
    'A customer technician, contractor or founder carries the sensors first.',
    'No cleanroom, ATEX area, valve operation, panel opening or automated maintenance decision.',
    'Timestamped checkpoint evidence, raw files and a human-approved exception record.',
    'Measured route time, completion rate, review minutes, false alerts and useful findings.',
    'A robot demonstration only after the customer accepts the evidence and economics.',
]:
    add_bullet(doc, text)

doc.add_heading('When to move to the next opportunity', level=1)
add_bullet(doc, 'Move to water-asset cleaning if buyers have clear confined-space or shutdown costs but will not pay for digital rounds.')
add_bullet(doc, 'Move to controlled-area UV-C only with a Quality or Validation sponsor and an OEM-supported qualification route.')
add_bullet(doc, 'Move to drainage records if industrial access is too slow but the property network produces an immediate paid survey.')
add_bullet(doc, 'Stop if two consecutive wedges produce only free-pilot interest or require custom hardware before payment.')

doc.add_heading('What not to do', level=1)
for text in [
    'Do not build humanoids, a custom chassis or a generic fleet platform.',
    'Do not buy solar, façade, cleaning or inspection robots because the demonstration looks impressive.',
    'Do not assume Pfizer access bypasses employer rules, Quality, procurement or validation.',
    'Do not travel to China until a paid route creates a precise hardware specification.',
]:
    add_bullet(doc, text)

final = doc.add_table(rows=1, cols=1)
cell = final.cell(0, 0)
set_cell_shading(cell, '1F4D46')
set_cell_margins(cell, top=145, start=180, bottom=145, end=180)
p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(0)
r = p.add_run('The company can become an autonomous operations layer for regulated facilities. The first business is one paid, observable route.')
r.bold = True
r.font.color.rgb = RGBColor(255, 255, 255)

# Metadata
props = doc.core_properties
props.title = 'The best robotics opportunities for us'
props.subject = 'Ireland-first robotics opportunity summary'
props.author = 'Sam Donworth'
props.keywords = 'robotics, Ireland, RaaS, pharma, industrial inspection'

doc.save(OUT)
print(OUT)
