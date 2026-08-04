from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/samdonworth/GroundZero/vault/briefs/2026-08-04-renewables-operations-intelligence-strategy.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.65)
sec.right_margin = Inches(0.65)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.2)
styles["Normal"].paragraph_format.space_after = Pt(3)
styles["Normal"].paragraph_format.line_spacing = 1.02
for name, size, color in [("Title", 25, "173047"), ("Heading 1", 16, "173047"), ("Heading 2", 11.5, "B38221")]:
    s = styles[name]
    s.font.name = "Aptos Display"
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(5)
    s.paragraph_format.space_after = Pt(3)

# Header/footer
header = sec.header.paragraphs[0]
header.text = "RENEWABLE INFRASTRUCTURE OPERATIONS INTELLIGENCE  |  4 AUGUST 2026"
header.style = styles["Normal"]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(7.5)
header.runs[0].font.color.rgb = RGBColor.from_string("5B6770")
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Confidential founder strategy  •  evidence-led validation only")
run.font.size = Pt(7.5)
run.font.color.rgb = RGBColor.from_string("5B6770")


def shade(p, fill="EEF2F4"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def title(t, sub=None):
    p = doc.add_paragraph(style="Title")
    p.add_run(t)
    if sub:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p2.add_run(sub)
        r.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor.from_string("B38221")


def h1(t): doc.add_heading(t, level=1)
def h2(t): doc.add_heading(t, level=2)

def body(t, bold_prefix=None, box=False):
    p = doc.add_paragraph()
    if bold_prefix and t.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(t[len(bold_prefix):])
    else:
        p.add_run(t)
    if box:
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.15)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(5)
        shade(p)
    return p


def bullets(items, level=0):
    for x in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Inches(0.2 + level * 0.18)
        p.paragraph_format.first_line_indent = Inches(-0.13)
        p.paragraph_format.space_after = Pt(1.5)
        p.add_run(x)


def page(): doc.add_page_break()

# PAGE 1

title("Renewable infrastructure operations intelligence", "Ireland-first solar operations, analytics and field assurance")
body("The business is credible, but not as a solar-cleaning company and not as another generic dashboard.", box=True)
h1("Verdict")
body("Start with an independent owner-side operations and assurance service for solar portfolios. Connect plant data, contractor work and field evidence so owners can find losses, prove action and run assets with fewer blind spots.")
h2("What the customer buys")
bullets([
    "Fewer unexplained generation and export losses.",
    "Faster triage across SCADA, inverter portals, meters and contractors.",
    "A traceable asset-to-fault-to-work-order-to-recovery history.",
    "Trusted monthly investor, lender and management reporting.",
    "Independent evidence that the O&M provider performed and the fault was resolved.",
])
h2("Why this is stronger than mowing first")
bullets([
    "Monitoring, reporting and contractor coordination recur through winter.",
    "Robotic mowing becomes one summer module in a year-round site contract.",
    "The asset is data, workflow and trust, rather than one seasonal machine.",
])
h2("The market is real")
bullets([
    "Ireland reached 2.7 GW connected solar by May 2026, including 1.59 GW utility-scale.",
    "Solar Ireland targets 8 GW by 2030 and estimates 0.8–1.3 GW annual additions may be needed.",
    "UK DESNZ reported 22.8 GW total solar by June 2026 and estimated about 59% was ground-mounted.",
    "Solar assets typically require 30+ years of operations and maintenance.",
])
h2("The constraint")
body("A part-time two-founder company cannot credibly offer full O&M on day one. Exclude 24/7 control, HV work, switching, guaranteed availability, spare-parts liability and turbine maintenance. Sell owner-side assurance during agreed hours, with specialist field partners.")

# PAGE 2
page()
title("The operating problem", "Owners have data and contractors, but often not one reliable operating picture")
h2("Fragmented inputs")
bullets([
    "SCADA, inverter and tracker portals; fiscal and export meters; weather and soiling sensors.",
    "EPC warranties, O&M ticketing, thermal-inspection PDFs, cleaning and vegetation contractors.",
    "PPA, market and curtailment data; planning, ecology, lender and insurer obligations.",
])
h2("Expensive failure modes")
bullets([
    "Inverter, string, communications or sensor failures noticed late.",
    "False alarms, poor baselines and inconsistent performance-ratio calculations.",
    "No clean handoff from alarm to action to verified recovery.",
    "Unclear responsibility across owner, EPC, OEM, O&M contractor and grid.",
    "Missing serial numbers, photographs and evidence that delay warranty claims.",
    "Vegetation shading, drainage, erosion, security and biodiversity issues tracked separately.",
])
h2("The five-step product")
bullets([
    "Signal: identify an anomaly from plant, meter, weather or image data.",
    "Diagnosis: estimate affected capacity, likely cause, urgency and value at risk.",
    "Action: route a work order to the responsible contractor or OEM.",
    "Evidence: capture readings, parts, photographs, cost and completion proof.",
    "Verification: confirm production recovery and close the loop.",
])
h2("Initial managed service")
bullets([
    "Read-only integrations and a verified asset register.",
    "Generation/export reconciliation and anomaly review.",
    "Owner-visible tickets with fault class, affected capacity and responsibility.",
    "O&M response, resolution, preventive-maintenance and repeat-fault tracking.",
    "Monthly operating and investor report with source-linked evidence.",
])
h2("AI that earns its keep")
bullets([
    "Normalise alarms across vendors; rank faults by lost-MWh exposure; identify repeats.",
    "Extract components and serial numbers; check contractor evidence; draft source-linked reports.",
    "No autonomous plant control and no AI-only electrical diagnosis presented as fact.",
])

# PAGE 3
page()
title("Technology and competition", "The innovation is integration and execution, not owning every tool")
h2("Cleaning and vegetation")
body("SolarPower Europe says cleaning frequency should be site-specific. Ireland's rainfall weakens blanket schedules. Measure soiling, rainfall, local fouling and expected recovery first. Cleaning robots and robotic mowers are deployment tools for suitable sites, not the company thesis.")
h2("Drones")
bullets([
    "Thermal/RGB inspection can localise faults quickly and inspect several MW per hour.",
    "Use at commissioning, after severe weather or when monitoring shows unexplained losses.",
    "Scheduled drone surveys are snapshots. Their value rises when findings are linked to continuous data and work orders.",
    "Partner first. A thermal drone alone does not create engineering, regulatory or insurance capability.",
])
h2("Planet and Google")
bullets([
    "PlanetScope at 3.7 m is useful for broad vegetation, flooding, access, construction and land-change monitoring, not module faults.",
    "SkySat at 50 cm adds detail but not thermal diagnosis and brings tasking, weather and cost constraints.",
    "Google Earth is context and history, not dependable operational cadence.",
    "Google Earth Engine public imagery is generally too coarse for panel defects; Google Solar API is primarily building/rooftop potential.",
])
h2("Incumbents")
bullets([
    "Irish service market: Galetech, EnergyPro, Optinergy, Clean Solar Solutions, Glás, SGR, Gener8 and Engineers With Drones.",
    "Software: Power Factors, GreenPowerMonitor, Clir, Raptor Maps and Sitemark.",
    "The gap is not absent technology. It is local deployment and owner accountability for portfolios that will not operate an enterprise stack well.",
])
h2("Wind")
body("Wind inspection and O&M are mature and higher-barrier. Owners care more about availability, blade erosion, lightning, component work, OEM oversight and safe repair than generic cleaning. Add wind reporting only through an existing customer, and use qualified inspection partners.")
h2("Installation robotics")
body("Utility construction robots are real, including AES Maximo, Terabase, Cosmic and Charge Robotics. Rooftop drone videos generally show lifting, not complete installation. Mounting, wind loads, weatherproofing, wiring, testing and certification remain. Track the sector; do not build or buy unproven hardware.")

# PAGE 4
page()
title("Economics and 90-day decision gate", "Payment and operational evidence before software or equipment")
h2("Scenario economics")
bullets([
    "Indicative mature-European utility OPEX of €7–€13/kWp/year implies an €11m–€21m Irish full-OPEX envelope. This is not all available to a newcomer.",
    "A hypothetical owner-assurance fee of €500–€2,000/MW/year implies an Ireland TAM of roughly €0.8m–€3.2m, and roughly €7.5m–€30m across Ireland plus estimated UK ground-mount.",
    "At 11% capacity factor and €70–€100/MWh, recovering 1% is worth about €6.7k–€9.6k/year on 10 MW, €33.7k–€48.2k on 50 MW and €67.5k–€96.4k on 100 MW.",
])
h2("Test offer")
bullets([
    "60–90 day operational baseline and leakage audit: test €4,000–€8,000 per site.",
    "Managed owner-side desk: test €750 monthly base plus €20/MW/month, with a minimum.",
    "Field work: transparent partner pass-through plus coordination or assurance fee.",
    "Five 30 MW sites under that formula produce about €81,000 ARR: validation income, not yet a two-founder company.",
])
h2("First customer")
bullets([
    "An independent Irish or UK owner with multiple sites or more than roughly 20 MW.",
    "Mixed portals or contractors, manual monthly reporting and no full internal asset-management team.",
    "Authority to grant read-only access and willingness to fund a paid diagnostic.",
])
h2("90-day sequence")
bullets([
    "Days 1–14: clear conflict/IP boundaries; interview 3–5 owners; inspect redacted workflows and monthly packs.",
    "Days 15–30: build a manual baseline with cleanly owned Evolv capabilities; price a paid pilot.",
    "Days 31–60: run read-only; measure hours removed, issues found, response delays and evidence gaps.",
    "Days 61–90: convert to a 12-month service. Build reusable software only after the same need appears across three customers.",
])
h2("Kill criteria")
bullets([
    "No owner will pay for a diagnostic; existing O&M already delivers trusted owner-side evidence; or the pain is only dashboard aesthetics.",
    "Unsafe control access, bespoke non-reusable integrations, mandatory 24/7 staffing or employment/IP conflict.",
])
body("NEXT GATE  •  One warm portfolio owner pays for an owner-side operational baseline and grants read-only data access.", box=True)

# Properties
props = doc.core_properties
props.title = "Renewable Infrastructure Operations Intelligence Strategy"
props.subject = "Ireland-first solar O&M, analytics, robotics and later wind"
props.author = "Hermes Agent for Sam Donworth"

doc.save(OUT)
print(OUT)
