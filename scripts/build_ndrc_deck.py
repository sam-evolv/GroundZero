from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUT = Path('/Users/samdonworth/Desktop/OpenHouse-NDRC-3-minute-deck.pptx')
OUT.parent.mkdir(parents=True, exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(8, 8, 8)
PANEL = RGBColor(18, 18, 18)
GOLD = RGBColor(212, 175, 55)
IVORY = RGBColor(244, 241, 234)
MUTED = RGBColor(178, 176, 168)
LINE = RGBColor(55, 52, 43)


def add_rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=18, color=IVORY, bold=False, font='Aptos', align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, tracking=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    if tracking is not None:
        r.font._element.set('spc', str(tracking))
    return box


def rich_lines(slide, lines, x, y, w, h):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, (content, size, color, bold, space_after) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space_after)
        r = p.add_run()
        r.text = content
        r.font.name = 'Aptos'
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return box


def base(slide, num, kicker):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    add_rect(slide, Inches(0.58), Inches(0.52), Inches(0.42), Inches(0.04), GOLD)
    text(slide, kicker.upper(), Inches(1.14), Inches(0.39), Inches(5.2), Inches(0.28), 10, GOLD, True)
    text(slide, f'0{num}', Inches(12.05), Inches(0.36), Inches(0.7), Inches(0.32), 10, MUTED, True, align=PP_ALIGN.RIGHT)
    add_rect(slide, Inches(0.58), Inches(7.02), Inches(12.17), Inches(0.012), LINE)
    text(slide, 'OPENHOUSE  |  NDRC PRE-ACCELERATOR', Inches(0.58), Inches(7.12), Inches(5.0), Inches(0.2), 8, MUTED, True)

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[6])
base(slide, 1, 'The problem')
text(slide, 'Every new home begins\nwith a loss of context.', Inches(0.62), Inches(1.1), Inches(7.7), Inches(1.65), 34, IVORY, True)
text(slide, 'At handover, vital home knowledge is scattered across PDFs, emails, warranties, commissioning records and support threads.', Inches(0.65), Inches(2.95), Inches(7.4), Inches(0.85), 17, MUTED)
add_rect(slide, Inches(8.92), Inches(1.12), Inches(3.78), Inches(4.88), PANEL, True, LINE)
text(slide, 'THE COST', Inches(9.28), Inches(1.50), Inches(2.5), Inches(0.28), 10, GOLD, True)
rich_lines(slide, [
    ('Homeowners inherit confusion.', 18, IVORY, True, 15),
    ('Developers inherit repeated aftercare, weak evidence and fragmented communication.', 15, MUTED, False, 25),
    ('Energy systems are installed, but not always understood or used with confidence.', 15, MUTED, False, 0),
], Inches(9.28), Inches(1.96), Inches(2.96), Inches(3.25))
text(slide, 'A home should not become harder to understand the day it is completed.', Inches(0.65), Inches(5.42), Inches(7.5), Inches(0.45), 15, GOLD, True)

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[6])
base(slide, 2, 'The product')
text(slide, 'OpenHouse turns handover\ninto a living home model.', Inches(0.62), Inches(1.1), Inches(8.0), Inches(1.55), 34, IVORY, True)
text(slide, 'We start where home context is created, then keep it useful through ownership.', Inches(0.65), Inches(2.86), Inches(7.3), Inches(0.45), 17, MUTED)
steps = [
    ('01', 'Approved home context', 'Documents, systems, warranties and handover evidence'),
    ('02', 'Homeowner clarity', 'A practical, home-specific source of help and guidance'),
    ('03', 'Developer intelligence', 'Aftercare signals, recurring issues and confidence gaps'),
]
for i, (n, head, sub) in enumerate(steps):
    x = Inches(0.65 + i * 4.1)
    add_rect(slide, x, Inches(4.03), Inches(3.62), Inches(1.42), PANEL, True, LINE)
    text(slide, n, x + Inches(0.26), Inches(4.29), Inches(0.4), Inches(0.2), 10, GOLD, True)
    text(slide, head, x + Inches(0.26), Inches(4.60), Inches(3.0), Inches(0.25), 15, IVORY, True)
    text(slide, sub, x + Inches(0.26), Inches(4.96), Inches(3.0), Inches(0.31), 10.5, MUTED)
text(slide, 'Not generic AI. A trusted, evidence-backed record of a specific home.', Inches(0.65), Inches(5.92), Inches(8.5), Inches(0.4), 15, GOLD, True)

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[6])
base(slide, 3, 'Proof and next step')
text(slide, 'Start with one scheme.\nProve the living home model. Expand.', Inches(0.62), Inches(1.1), Inches(8.2), Inches(1.55), 31, IVORY, True)
add_rect(slide, Inches(0.65), Inches(3.1), Inches(5.72), Inches(2.46), PANEL, True, LINE)
text(slide, 'OPERATIONAL PROOF', Inches(0.96), Inches(3.43), Inches(2.6), Inches(0.25), 10, GOLD, True)
rich_lines(slide, [
    ('Built from inside residential delivery.', 18, IVORY, True, 13),
    ('Live product with real scheme data across four Longview Estates schemes.', 15, MUTED, False, 0),
], Inches(0.96), Inches(3.87), Inches(4.88), Inches(1.18))
add_rect(slide, Inches(6.72), Inches(3.1), Inches(5.98), Inches(2.46), PANEL, True, LINE)
text(slide, 'WHAT NDRC UNLOCKS', Inches(7.03), Inches(3.43), Inches(3.0), Inches(0.25), 10, GOLD, True)
rich_lines(slide, [
    ('Turn operational proof into an externally repeatable business.', 18, IVORY, True, 13),
    ('Sharpen the first paid wedge, validate it with outside developers, and build the credibility for the next funding stage.', 15, MUTED, False, 0),
], Inches(7.03), Inches(3.87), Inches(5.1), Inches(1.18))
text(slide, 'The ambition: every home has a useful, trusted intelligence layer from day one.', Inches(0.65), Inches(6.0), Inches(11.8), Inches(0.42), 16, GOLD, True)

prs.save(OUT)
print(OUT)
